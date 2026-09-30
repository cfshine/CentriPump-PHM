"""文档处理模块：加载 + 预处理（清洗 / 结构化 / 切片）。

- 加载：把 PDF / Excel / Markdown 解析成 LangChain Document（CSV 为旁路，
  需显式传入单个 .csv 文件路径才会解析，见 load_documents）。
- 预处理：clean_documents（清洗）→ structure_documents（结构化，出 DocRecord）
  → split_documents（切片，出 Chunk）。

本模块不含向量化与向量库逻辑，那些在 knowledge_base.py。
"""

import hashlib
import json
import re

from collections import Counter
from html import unescape as html_unescape
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pydantic import BaseModel, Field

# 支持加载的文件后缀 -> 对应加载方式
PDF_SUFFIXES = {".pdf"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
MARKDOWN_SUFFIXES = {".md", ".markdown"}
CSV_SUFFIXES = {".csv"}
SUPPORTED_SUFFIXES = PDF_SUFFIXES | EXCEL_SUFFIXES | MARKDOWN_SUFFIXES

# CSV 分段列：按此顺序取表头中第一个命中的列，把连续相同取值的行合并成一个 Document
CSV_SEGMENT_COLUMNS = ("operating_state", "operating_condition", "status", "state", "alarm_code")
# 单个 CSV Document 最多容纳的行数（也是表头中没有分段列时的分块大小）
CSV_ROWS_PER_DOC = 50

# -------------------- 清洗：页眉页脚剔除 --------------------
HEADER_WINDOW = 3          # 每页取首 / 末 K 行作为候选池
HEADER_MIN_RATIO = 0.6     # 候选行出现率 >= 此比例才判为页眉页脚
HEADER_MAX_LEN = 60        # 候选行长度上限：正文长句即使重复也不该删
MIN_PAGES_FOR_HEADER = 3   # PDF 页数少于此值不做统计（样本太少，容易误删正文）

# -------------------- 结构化 --------------------
MIN_RECORD_CHARS = 8       # 短于此长度的块丢弃（页码残留、孤立符号等噪声）
# 连续 >= 这么多条「同层级、中间无正文」的标题判为目录段，不作为章节起点
TOC_RUN_MIN = 3

# -------------------- 切片 --------------------
CHUNK_SIZE = 500           # 与 manual_agent/retriever.py 默认值对齐
CHUNK_OVERLAP = 80
MAX_TABLE_CHARS = CHUNK_SIZE * 2   # 表格超过此长度才按行组切，否则整表一块
# 中文分隔符：英文默认分隔符（\n\n \n 空格）会把中文长句从中间截断
CHINESE_SEPARATORS = ["\n\n", "\n", "。", "！", "？", "；", "，", " ", ""]

ZERO_WIDTH_CHARS = re.compile(r"[\ufeff\u200b\u200c\u200d]")
MARKDOWN_ESCAPE = re.compile(r"\\([_*#\[\]()~`])")
MD_HEADING = re.compile(r"^(#{1,6})\s+(\S.*)$")
CN_HEADING = re.compile(r"^[一二三四五六七八九十]+[、.．]\s*\S")
NUM_HEADING = re.compile(r"^(\d+(?:\.\d+)*)[、.．]?\s+(\S.*)$")
# 附录标题：`附录B SCADA 监测点定义` / `附录 E 仿真运行状态定义` / `Appendix A ...`。
# 字母必须紧跟其后、且**后面是空白或行尾**——否则会误伤正文里的「附录B中列出了…」。
APPENDIX_HEADING = re.compile(r"^(?:附录|Appendix)\s*[A-Za-z](?:\s|$)")
# 「数字 + 单位」开头的行是数据行（如 "15 min 短期趋势"），不是标题
UNIT_TOKEN = re.compile(
    r"^(?:min|h|s|ms|Hz|rpm|r/min|°C|℃|mm|cm|km|m|MPa|kPa|bar|kW|W|V|A|mA|%|"
    r"L/min|m3/h|L|kg|N·m|dB)\b"
)
LIST_ITEM = re.compile(r"^\s*(?:[-*+]|\d+[.、)])\s+\S")
TABLE_ROW = re.compile(r"^\s*\|")
TABLE_SEPARATOR = re.compile(r"^\s*\|[\s:|-]*\|\s*$")


class DocRecord(BaseModel):
    """结构化的语义块：章节片段 / 整张表格。阶段 3 的产物，见设计文档 4.2。"""

    doc_id: str
    source: str
    file_type: str
    section_path: List[str] = []
    content_type: str = "text"        # "text" | "table" | "list"
    locator: Dict[str, Any] = {}
    raw_text: str
    char_count: int


class Chunk(BaseModel):
    """检索的最小单元，切片阶段（4）的产物，见设计文档 4.3。

    不存向量：向量化由 `add_documents` 内部经 `embedding_function` 完成，
    没有哪一步会把向量回填到这里，留着这个字段只会误导读代码的人。
    """

    chunk_id: str
    doc_id: str
    text: str
    chunk_index: int
    source: str
    file_type: str
    section_path: List[str] = []
    content_type: str = "text"
    locator: Dict[str, Any] = {}
    char_count: int


class LoadStats(BaseModel):
    """加载阶段的统计回报。

    加载阶段原来只打印日志、不返回结构化结果，导致「少了哪份资料」无法被上游感知，
    `BuildReport` 的 skipped / failed 也就填不出来（见设计文档 4.6）。

    用法是**出参**而不是改返回值，这样 `load_documents` 的返回契约保持不变：

        stats = LoadStats()
        documents = load_documents("docs", stats=stats)

    注意区分两个计数：
    - `files_scanned` 是**原始文件数**（`docs/` 里 2 个文件）；
    - `docs_loaded` 是**Document 数**——而 Document 的粒度是一页 / 一个 Sheet / 一个文件，
      同一批资料这两个数通常差很多（实测 2 个文件 -> 24 个 Document），混为一谈会误导。
    """

    files_scanned: int = 0
    docs_loaded: int = 0
    skipped: List[str] = Field(default_factory=list)   # 不支持的后缀
    failed: List[str] = Field(default_factory=list)    # 读取失败的文件及原因


# -------------------- 文档加载 --------------------
def load_documents(path: Union[str, Path], stats: Optional[LoadStats] = None) -> List[Document]:
    """加载 PDF / Excel / Markdown 文档，返回 LangChain Document 列表。

    path 既可以是一个文件，也可以是一个目录；传目录时会递归遍历，
    只收集后缀在 SUPPORTED_SUFFIXES 中的文件（不含 .csv，见下）。
    单个文件读取失败只打印警告并跳过，不影响其余文件。

    `stats` 是可选出参：传入 `LoadStats()` 即可拿到「扫了几个文件、出了几个 Document、
    跳过/失败哪些」的结构化结果，供上游 `BuildReport` 汇总。**注意区分文件数与
    Document 数**——Document 的粒度是「PDF 一页 / Excel 一个 Sheet / Markdown 整个文件」，
    实测 `docs/` 下 2 个文件会产出 24 个 Document。

    CSV 是旁路：直接传某个 .csv 文件路径仍会解析（走 _load_csv），
    但目录遍历不会收集 CSV，因此工况数据不会进入知识库链路。

    ⚠️ 入参会先规整成绝对路径：`metadata["source"]` 取自文件路径，而它是确定性 id
    的输入之一。若不做规整，同一份资料用 `docs/a.pdf` 与 `/abs/path/docs/a.pdf`
    两种写法构建，会算出两套 id → 重复构建不覆盖而是**新增一份**，幂等静默失效。

    依赖：
        pip install pypdf openpyxl
    """
    path = Path(path).resolve()   # 统一路径写法，保证 id 稳定（见上）
    if not path.exists():
        raise FileNotFoundError(f"路径不存在：{path}")

    if path.is_file():
        files = [path]
    else:
        files = sorted(
            p for p in path.rglob("*")
            if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
        )

    documents: List[Document] = []
    for file in files:
        if stats is not None:
            stats.files_scanned += 1     # 扫描数含后面失败/跳过的文件，与上方日志口径一致
        suffix = file.suffix.lower()
        try:
            if suffix in PDF_SUFFIXES:
                docs = _load_pdf(file)
            elif suffix in EXCEL_SUFFIXES:
                docs = _load_excel(file)
            elif suffix in MARKDOWN_SUFFIXES:
                docs = _load_markdown(file)
            elif suffix in CSV_SUFFIXES:
                docs = _load_csv(file)
            else:
                print(f"[WARN] 不支持的文件类型，已跳过：{file}")
                if stats is not None:
                    stats.skipped.append(str(file))
                continue
        except Exception as exc:  # noqa: BLE001
            print(f"[WARN] 加载失败 {file}: {exc}")
            if stats is not None:
                stats.failed.append(f"{file}: {exc}")
            continue

        documents.extend(docs)
        if stats is not None:
            stats.docs_loaded += len(docs)
        print(f"[INFO] 已加载 {file.name} -> {len(docs)} 个 Document")

    print(
        f"[INFO] 本次共扫描 {len(files)} 个文件 -> {len(documents)} 个 Document"
        f"（PDF 每页 / Excel 每 Sheet / Markdown 整个文件 各算 1 个）"
    )
    return documents


def _load_pdf(file: Path) -> List[Document]:
    """按页加载 PDF，每页一个 Document。"""
    try:
        from langchain_community.document_loaders import PyPDFLoader
    except ImportError as exc:
        raise ImportError("加载 PDF 需要 pypdf，请先执行：pip install pypdf") from exc

    docs = PyPDFLoader(str(file)).load()
    for doc in docs:
        doc.metadata["source"] = str(file)
        doc.metadata["file_type"] = "pdf"
    return docs


def _load_excel(file: Path) -> List[Document]:
    """按 Sheet 加载 Excel，每个 Sheet 一个 Document，单元格以 " | " 连接成行。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("加载 Excel 需要 openpyxl，请先执行：pip install openpyxl") from exc

    # data_only=True 读取公式的计算结果而非公式本身；read_only 降低大表内存占用
    workbook = load_workbook(str(file), data_only=True, read_only=True)
    documents: List[Document] = []
    try:
        for sheet in workbook.worksheets:
            lines: List[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if cell is None else str(cell).strip() for cell in row]
                if not any(cells):  # 跳过空行
                    continue
                lines.append(" | ".join(cells).rstrip(" |"))
            if not lines:
                continue
            documents.append(
                Document(
                    page_content="\n".join(lines),
                    metadata={
                        "source": str(file),
                        "file_type": "excel",
                        "sheet_name": sheet.title,
                    },
                )
            )
    finally:
        workbook.close()
    return documents


def _load_markdown(file: Path) -> List[Document]:
    """加载 Markdown，整个文件作为一个 Document（保留标题等 Markdown 结构）。"""
    text = file.read_text(encoding="utf-8", errors="ignore")
    if not text.strip():
        return []
    return [
        Document(
            page_content=text,
            metadata={"source": str(file), "file_type": "markdown"},
        )
    ]


def _load_csv(file: Path) -> List[Document]:
    """按「工况分段」加载 CSV，每个 Document 首行是表头，其后是该段的数据行。

    分段依据是表头中第一个命中的 CSV_SEGMENT_COLUMNS 列（如 operating_state）：
    连续取值相同的行合并成一个 Document。这样一段连续的工况/报警会成为语义完整的
    检索单元，而不是让整份遥测数据被切片器从中间随机截断。若表头没有这类列，
    则退化为按 CSV_ROWS_PER_DOC 行固定分块。两种情况下单块行数都以
    CSV_ROWS_PER_DOC 为上限。

    只依赖标准库 csv，不引入额外依赖。中文 CSV 常见 GBK/GB18030 编码，
    这里按 utf-8-sig -> gbk -> gb18030 依次尝试。
    """
    import csv
    import io

    content = None
    for encoding in ("utf-8-sig", "gbk", "gb18030"):
        try:
            content = file.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if content is None:
        raise ValueError(f"无法识别文件编码（已尝试 utf-8-sig / gbk / gb18030）：{file}")

    rows = [row for row in csv.reader(io.StringIO(content)) if any(c.strip() for c in row)]
    if len(rows) < 2:  # 只有表头或空文件
        return []

    header, data_rows = rows[0], rows[1:]
    header_line = " | ".join(header)

    lower_header = [h.strip().lower() for h in header]
    segment_idx = next(
        (lower_header.index(name) for name in CSV_SEGMENT_COLUMNS if name in lower_header),
        None,
    )

    documents: List[Document] = []
    block: List[List[str]] = []
    block_key = ""

    def flush() -> None:
        """把当前累积的行落成一个 Document。"""
        if not block:
            return
        metadata = {
            "source": str(file),
            "file_type": "csv",
            "row_count": len(block),
        }
        if segment_idx is not None:
            metadata["segment_column"] = header[segment_idx]
            metadata["segment_value"] = block_key
        documents.append(
            Document(
                page_content="\n".join([header_line] + [" | ".join(row) for row in block]),
                metadata=metadata,
            )
        )
        block.clear()

    for row in data_rows:
        # 短行（字段数不足）归入空键，不至于下标越界
        key = row[segment_idx].strip() if segment_idx is not None and segment_idx < len(row) else ""
        if block and (key != block_key or len(block) >= CSV_ROWS_PER_DOC):
            flush()
        block_key = key
        block.append(row)
    flush()

    return documents


# -------------------- 文档清洗与处理 --------------------
def clean_documents(documents: List[Document]) -> List[Document]:
    """跨文档清洗：去页眉页脚 / 去转义 / 归一化空白。

    签名是「列表 → 列表」而非「单文档 → 单文档」——页眉页脚判定依赖全量统计，
    只看一页无从知道某行是不是重复页眉。规则见《RAG类设计与方法结构.md》5.2。

    返回**新的 Document 列表**，数量与顺序不变，只改 `page_content` 与
    `metadata["cleaning"]`；原始对象不动，便于对比排查。
    """
    documents = [
        Document(page_content=_normalize_text(doc.page_content), metadata=dict(doc.metadata))
        for doc in documents
    ]

    header_lines = _detect_header_lines(documents)
    if not header_lines:
        return documents

    cleaned: List[Document] = []
    for doc in documents:
        lines = doc.page_content.split("\n")
        removed = sorted({line.strip() for line in lines if line.strip() in header_lines})
        if not removed:
            cleaned.append(doc)
            continue
        kept = [line for line in lines if line.strip() not in header_lines]
        cleaned.append(
            Document(
                page_content="\n".join(kept).strip(),
                metadata={**doc.metadata, "cleaning": {"removed_header_lines": removed}},
            )
        )
    return cleaned


def _normalize_text(text: str) -> str:
    """逐文档的规范化：去 BOM/零宽字符、还原 HTML 实体、去 Markdown 转义、归一空白。

    规则顺序与设计文档 5.2 一致（1/2/3 → 5/6/7/8）；页眉页脚剔除（规则 4）因为要
    跨文档统计，放在 `_detect_header_lines` 里单独做——且在归一化之后做，这样各页
    之间是拿归一化后的行比对，重复行更容易对齐。
    """
    text = ZERO_WIDTH_CHARS.sub("", text)
    text = html_unescape(text).replace("\xa0", " ")   # &nbsp; -> 普通空格
    text = MARKDOWN_ESCAPE.sub(r"\1", text)           # T\_bearing -> T_bearing
    text = re.sub(r"[ \t]+$", "", text, flags=re.M)   # 行尾硬换行空格
    text = re.sub(r"[^\S\n]+", " ", text)             # 行内连续空白折叠
    text = re.sub(r"\n{3,}", "\n\n", text)            # 连续空行折叠
    return text.strip()


def _detect_header_lines(documents: List[Document]) -> set:
    """跨文档统计重复的页眉页脚行，返回应删除的行集合。

    只对 PDF 生效（Excel / Markdown 无页眉页脚概念）。判定条件三者同时满足：
    出现于页面首 / 末 K 行窗口内、出现率 >= HEADER_MIN_RATIO、长度 <= HEADER_MAX_LEN。
    只用频率不行——实测该语料页脚页页不同，必须叠加「位置」这一条件。
    """
    pdf_docs = [d for d in documents if d.metadata.get("file_type") == "pdf"]
    if len(pdf_docs) < MIN_PAGES_FOR_HEADER:
        return set()

    counts: Counter = Counter()
    for doc in pdf_docs:
        lines = [line.strip() for line in doc.page_content.split("\n")]
        window = set(lines[:HEADER_WINDOW] + lines[-HEADER_WINDOW:])
        for line in window:      # 同一页内重复出现只算一次
            if line:
                counts[line] += 1

    threshold = HEADER_MIN_RATIO * len(pdf_docs)
    return {
        line for line, count in counts.items()
        if count >= threshold and len(line) <= HEADER_MAX_LEN
    }


# -------------------- 结构化 --------------------
def structure_documents(documents: List[Document]) -> List[DocRecord]:
    """解析标题层级与块类型，把每份文档拆成带章节路径的语义块（DocRecord）。

    三步：按行维护 section_path 栈（标题判级见 `_heading_level`）→ 识别块类型
    （表格 / 列表 / 正文）→ 「章节变化」或「块类型变化」处切断。见设计文档 5.3。

    两个实现要点：
    - **章节栈跨页延续**：PDF 是「一页一个 Document」，章节会跨页，所以栈按
      `source` 分组维护，而不是每个 Document 重置（否则第 4 页起就没有章节归属了）。
    - **表格整体产出**：整张表格是一个块，不与前后正文混在一起，否则后续切片
      会把表切断。

    出口处再走一道 `_drop_front_matter`：丢掉首个标题之前的封面 / 目录 / 文件控制表。
    """
    records: List[DocRecord] = []
    stack: List[str] = []
    current_source: Optional[str] = None
    ordinal = 0            # 同一文件内的块序号，用于生成稳定 doc_id

    for doc in documents:
        source = doc.metadata.get("source", "")
        if source != current_source:
            stack = []     # 换文件，章节栈重置
            current_source = source
            ordinal = 0

        lines = doc.page_content.split("\n")
        headings = _heading_indices(lines)
        table_rows = _table_line_indices(lines)

        block: List[str] = []
        block_type = ""
        block_section: List[str] = []

        def flush() -> None:
            """把当前累积的行落成一个 DocRecord。"""
            nonlocal block, block_type, ordinal
            if block:
                ordinal += 1
                raw_text = "\n".join(block)
                if len(raw_text) >= MIN_RECORD_CHARS:
                    records.append(
                        DocRecord(
                            doc_id=_doc_id(source, locator, ordinal),
                            source=source,
                            file_type=doc.metadata.get("file_type", ""),
                            section_path=list(block_section),
                            content_type=block_type,
                            locator=locator,
                            raw_text=raw_text,
                            char_count=len(raw_text),
                        )
                    )
            block = []
            block_type = ""

        locator = _locator(doc.metadata)

        for index, line in enumerate(lines):
            stripped = line.strip()
            if not stripped:
                continue

            level = headings.get(index)
            if level is not None:
                flush()                       # 标题一定另起一块
                while len(stack) >= level:    # 退到父层级
                    stack.pop()
                stack.append(_clean_title(stripped))
                block_section = list(stack)
                block_type = "text"
                block.append(stripped)        # 标题文本并入下一块，检索时自带语境
                continue

            line_type = _line_type(index, lines, table_rows)
            if block and (line_type != block_type or list(stack) != block_section):
                flush()
            if not block:
                block_section = list(stack)
            block_type = line_type
            block.append(stripped)

        flush()

    return _drop_front_matter(records)


def _drop_front_matter(records: List[DocRecord]) -> List[DocRecord]:
    """丢掉「首个标题之前」的前置页：封面 / 目录 / 文件控制表。

    **为什么必须丢**（2026-09-30 实测）：这些块是**实体词最密集**的段落——封面和文件控制页
    反复出现「HCP-80-50-200A」「离心泵组」这类词，却永远不含运维答案。后果有两条：

    - 向量召回时它们的余弦分并不低（实测占第 2、3 位）；
    - **重排会把它们顶到最前**。cross-encoder 对词汇重叠的奖励压过了「是否真能回答问题」，
      实测有一条查询（报警代码/阈值）唯一正确的片段 `8.5 仿真报警代码` 因此被挤出 top3，
      换成封面与文件控制页。

    **判定依据是结构而不是关键词**：`section_path` 为空 ⟺ 该块在本文档的首个标题之前。
    这条等价关系来自 `structure_documents` 维护章节栈的方式——栈只在遇到标题时追加，
    所以首个标题之前的块拿不到任何路径。用结构判定比匹配「封面」「目录」这类字样稳，
    换一份资料不用改规则，也不怕正文明文里出现「目录」二字。

    ⚠️ **通篇没有标题的文档必须整份保留**：那种情况下所有块都没有 `section_path`，
    照「空路径就丢」会把它整份清空。所以先看这份文档有没有出现过标题，出现过才丢。

    ⚠️ 这是在**全部 ordinal 分配完之后**做过滤，不会改动幸存块的 `doc_id`——但库里
    被丢掉的旧块**不会自动消失**（id 对不上，upsert 覆盖不到），改完必须带
    `reset=True` 重建一次（见设计文档 5.8）。
    """
    has_heading: Dict[str, bool] = {}
    for record in records:
        if record.section_path:
            has_heading[record.source] = True

    return [
        record
        for record in records
        if record.section_path or not has_heading.get(record.source)
    ]


def _line_type(index: int, lines: List[str], table_rows: set) -> str:
    """判定单行属于哪种块类型。"""
    if index in table_rows:
        return "table"
    if LIST_ITEM.match(lines[index].strip()):
        return "list"
    return "text"


def _table_line_indices(lines: List[str]) -> set:
    """找出「连续 | 行」构成的表格块行号。

    只认带分隔行（`| :--- |`）的连续管道行，避免正文里偶发的单个 `|` 被误判成表格。
    """
    indices: set = set()
    run: List[int] = []
    for index, line in enumerate(lines + [""]):     # 多压一个空行，收尾时能触发结算
        if TABLE_ROW.match(line.strip()):
            run.append(index)
            continue
        if run:
            if any(TABLE_SEPARATOR.match(lines[i].strip()) for i in run):
                indices.update(run)
            run = []
    return indices


def _heading_level(line: str) -> Optional[int]:
    """判断一行是不是标题，是则返回层级（1 起），否则返回 None。

    四类规则按「可信度从高到低」依次尝试：Markdown `#` → 数字编号 → 附录 → 中文编号。

    PDF 抽出来的纯文本没有加粗信息，只能用「编号 + 行长」兜底。实测该语料 PDF 的标题
    形如 `1 安全使用须知` / `2.1 设备定位`，但也有大量形似的假标题（表格行
    `100 45 74 3.3 正常曲线基准`、清单项 `1 泵体充液 … □`、数据行 `15 min 短期趋势…`），
    所以额外要求：剩余部分不能还含数字、不能以单位开头、不能以句读结尾，且裸数字编号
    只认很短的短行。

    ⚠️ **附录规则是后补的**（2026-09-30）：此前只认数字/中文编号，`附录B SCADA 监测点定义`
    一个都不匹配，于是章节栈**停在上一个标题** `11.3 维修后基线确认` 不再更新，
    导致附录 B/C/D/E 的内容（共 16 条 chunk、占全库 16%）全被挂到那个小节点名下——
    喂给 LLM 的出处是错的。这与 `_heading_indices` 里已处理的「目录段让栈停在最后一条
    目录项」是同一类问题的不同入口。
    """
    match = MD_HEADING.match(line)
    if match:
        return len(match.group(1))

    if len(line) > 30:
        return None

    match = NUM_HEADING.match(line)
    if match:
        number, rest = match.group(1), match.group(2)
        if re.search(r"\d", rest):
            return None
        if UNIT_TOKEN.match(rest):
            return None
        if "□" in rest or rest.endswith(("。", "，", "；", "：")):
            return None
        if "." not in number and len(line) > 20:      # 裸数字编号只认短行
            return None
        return number.count(".") + 1

    # 附录与中文编号同为一级：`附录B SCADA 监测点定义` 是和 `1 安全使用须知` 平级的分部。
    # 放在数字规则之后，是为了让 `1.2 标准运行参数基线` 这类先被更具体的规则认走。
    if (
        APPENDIX_HEADING.match(line)
        and not line.endswith(("。", "，", "；", "："))
    ):
        return 1

    if CN_HEADING.match(line) and not line.endswith(("。", "，", "；", "：")):
        return 1
    return None


def _heading_indices(lines: List[str]) -> Dict[int, int]:
    """返回 {行号: 层级}，并过滤掉目录段。

    实测该语料 PDF 第 1 页有一段「目录」：连续 11 条一级标题、中间没有任何正文。
    不处理的话会多出 11 个只有标题的碎块，而且章节栈会停在最后一个目录项上，把随后
    的正文错挂到「11 操作记录与维修记录」下。

    判定条件：连续 >= TOC_RUN_MIN 条**同层级**标题且中间无正文。要求「同层级」是为了
    不误伤合法的父子标题嵌套（实测 `2 设备概述与型号说明` 紧跟 `2.1 设备定位`）。

    只对**没有 Markdown 标题**的文档做这一步——`#` 是作者显式写的，可信；PDF 抽出的
    文本里标题全靠启发式判断，才需要额外佐证。
    """
    candidates: Dict[int, int] = {}
    has_explicit_heading = False
    for index, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            continue
        if MD_HEADING.match(stripped):
            has_explicit_heading = True
        level = _heading_level(stripped)
        if level is not None:
            candidates[index] = level

    if has_explicit_heading or not candidates:
        return candidates

    kept: Dict[int, int] = {}
    run: List[int] = []

    def settle() -> None:
        if len(run) < TOC_RUN_MIN:
            kept.update({index: candidates[index] for index in run})
        elif run:
            print(f"[INFO] 疑似目录段，跳过 {len(run)} 条标题：" f"{lines[run[0]].strip()[:20]} …")

    for index in sorted(candidates):
        if run and candidates[index] == candidates[run[-1]] and _no_body_between(lines, run[-1], index):
            run.append(index)
        else:
            settle()
            run = [index]
    settle()
    return kept


def _no_body_between(lines: List[str], start: int, end: int) -> bool:
    """判断两行之间是否没有正文（只有空行）。"""
    return not any(line.strip() for line in lines[start + 1 : end])


def _clean_title(line: str) -> str:
    """把标题行压成纯文本：去掉 `#` 与 `**加粗**` 之类的 Markdown 装饰。"""
    title = MD_HEADING.sub(r"\2", line)
    return title.strip().strip("*_`").strip()


def _locator(metadata: dict) -> Dict[str, Any]:
    """从 Document 元数据里提取溯源定位信息（页码 / Sheet 名）。"""
    locator: Dict[str, Any] = {}
    if "page" in metadata:
        locator["page"] = metadata["page"]
    if "sheet_name" in metadata:
        locator["sheet_name"] = metadata["sheet_name"]
    return locator


def _doc_id(source: str, locator: Dict[str, Any], ordinal: int) -> str:
    """确定性块 id。

    设计文档 4.4 给的是 `md5(source + 定位信息)`，那是「一页 = 一条」的口径；这里
    一页会产出多个块，所以把块序号也纳入哈希，否则同页的块会撞 id、互相覆盖。
    """
    raw = f"{source}\x1f{json.dumps(locator, sort_keys=True, ensure_ascii=False)}\x1f{ordinal}"
    return hashlib.md5(raw.encode("utf-8")).hexdigest()


# -------------------- 文档切片 --------------------
def split_documents(records: List[DocRecord]) -> List[Chunk]:
    """按块类型切片：结构优先、长度兜底，输出 Chunk（设计文档 5.4）。

    - `table`：整表一块；超过 MAX_TABLE_CHARS 才按行组切，且每组重复表头行；
    - `list` ：整块保留，超长时只在列表项边界切，不切断单个列表项；
    - `text` ：用中文分隔符的 RecursiveCharacterTextSplitter。

    `section_path` 只保留在 Chunk 字段里、不注入正文（设计文档决策点：注入会污染
    「按原文显示引用」，且重复占用 embedding 预算）。
    """
    if not records:
        return []

    chunks: List[Chunk] = []
    for record in records:
        if record.content_type == "table":
            texts = _split_table(record.raw_text)
        elif record.content_type == "list":
            texts = _split_list(record.raw_text)
        else:
            texts = _split_text(record.raw_text)

        for index, text in enumerate(texts):
            chunks.append(
                Chunk(
                    chunk_id=_chunk_id(record.doc_id, index, text),
                    doc_id=record.doc_id,
                    text=text,
                    chunk_index=index,
                    source=record.source,
                    file_type=record.file_type,
                    section_path=list(record.section_path),
                    content_type=record.content_type,
                    locator=dict(record.locator),
                    char_count=len(text),
                )
            )
    return chunks


def _split_text(text: str) -> List[str]:
    """正文切片：递归字符切分，中文标点优先。"""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=CHINESE_SEPARATORS,
    )
    return [piece for piece in splitter.split_text(text) if piece.strip()]


def _split_table(text: str) -> List[str]:
    """表格切片：整表优先，超长时按行组切并给每组补回表头。"""
    if len(text) <= MAX_TABLE_CHARS:
        return [text]

    lines = text.split("\n")
    prefix = [lines[0]]
    body_start = 1
    if len(lines) > 1 and TABLE_SEPARATOR.match(lines[1].strip()):
        prefix.append(lines[1])          # 分隔行也要跟着表头一起重复
        body_start = 2

    groups: List[str] = []
    current: List[str] = []
    size = 0
    for line in lines[body_start:]:
        if current and size + len(line) > CHUNK_SIZE:
            groups.append("\n".join(prefix + current))
            current, size = [], 0
        current.append(line)
        size += len(line)
    if current:
        groups.append("\n".join(prefix + current))
    return groups


def _split_list(text: str) -> List[str]:
    """列表切片：只在列表项边界切分，不切断单个列表项。"""
    if len(text) <= CHUNK_SIZE:
        return [text]

    items: List[str] = []
    for line in text.split("\n"):
        if LIST_ITEM.match(line) or not items:
            items.append(line)           # 新列表项（含其后的续行）
        else:
            items[-1] = f"{items[-1]}\n{line}"

    groups: List[str] = []
    current: List[str] = []
    size = 0
    for item in items:
        if current and size + len(item) > CHUNK_SIZE:
            groups.append("\n".join(current))
            current, size = [], 0
        current.append(item)
        size += len(item)
    if current:
        groups.append("\n".join(current))

    # 单个列表项本身就超长 → 只能退回按长度切
    return [piece for group in groups for piece in _split_text(group)]


def _chunk_id(doc_id: str, index: int, text: str) -> str:
    """确定性 chunk id（设计文档 4.4）：同一内容重复构建得到同一个 id，保证幂等。"""
    raw = f"{doc_id}::{index}::{text[:64]}"
    return hashlib.md5(raw.encode("utf-8")).hexdigest()
